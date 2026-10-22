"""Exercises in the plan, how-to videos, and progress from manual workouts."""
import json

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(fake_llm):
    import main
    from milo.coach import Coach
    from milo.storage import Storage

    main.store = Storage(":memory:")
    main.coach = Coach(main.store)
    return TestClient(main.app)


def save_plan(client, fake_llm, plan):
    fake_llm.replies = [f"Plan.\n```json\n{json.dumps(plan)}\n```"]
    assert client.post("/api/coaching/message", json={"text": "plan please"}).json()["plan_saved"]


def test_exercises_are_kept_and_optional(client, fake_llm, plan_json):
    from milo.schemas import PlannedSession

    assert PlannedSession(day="Monday", type="run").exercises == []
    save_plan(client, fake_llm, plan_json)
    week1 = client.get("/api/plan").json()["plan"]["weeks"][0]["sessions"]
    strength = next(s for s in week1 if s["type"] == "strength")
    assert strength["exercises"][0]["name"] and strength["exercises"][0]["video_query"]


def test_downloads_list_exercises(client, fake_llm, plan_json):
    from openpyxl import load_workbook
    import io

    save_plan(client, fake_llm, plan_json)
    wb = load_workbook(io.BytesIO(client.get("/api/plan/export.xlsx").content))
    notes = " ".join(str(r[8].value or "") for r in wb["Plan"].iter_rows(min_row=5))
    assert "Bodyweight squat 3 x 12" in notes


def test_manual_workout_in_a_future_week_counts(client, fake_llm, plan_json):
    """Bug report: with the clock pinned (MILO_TODAY), Progress ignored manual workouts."""
    save_plan(client, fake_llm, plan_json)
    before = {w["week_number"]: w for w in client.get("/api/progress").json()}
    assert before[4]["status"] == "upcoming"
    client.post("/api/activities/manual", json={"date": "2026-11-03", "type": "run", "duration_min": 35})
    after = {w["week_number"]: w for w in client.get("/api/progress").json()}
    assert after[4]["status"] != "upcoming" and after[4]["actual_duration_min"] == 35


def test_videos_without_key_give_a_search_link(client, monkeypatch):
    from milo import config

    monkeypatch.setattr(config, "YOUTUBE_API_KEY", "")
    r = client.get("/api/videos", params={"q": "Goblet  Squat form"}).json()
    assert r["source"] == "link" and r["items"] == []
    assert r["search_url"].startswith("https://www.youtube.com/results?search_query=goblet+squat+form")


def test_videos_with_key_are_cached(client, monkeypatch):
    from milo import config, videos

    calls = []

    def fake(q, key, limit):
        calls.append(q)
        return [{"id": "abc123", "title": "Squat &amp; form", "channel": "Coach", "thumbnail": ""}]

    monkeypatch.setattr(config, "YOUTUBE_API_KEY", "test-key")
    monkeypatch.setattr(videos, "fetcher", fake)
    for _ in range(3):
        r = client.get("/api/videos", params={"q": "squat form"}).json()
    assert r["source"] == "youtube" and r["items"][0]["id"] == "abc123"
    assert calls == ["squat form"]  # one search, then the cache


def test_videos_fall_back_when_youtube_fails(client, monkeypatch):
    from milo import config, videos

    def broken(q, key, limit):
        raise OSError("quotaExceeded")

    monkeypatch.setattr(config, "YOUTUBE_API_KEY", "test-key")
    monkeypatch.setattr(videos, "fetcher", broken)
    r = client.get("/api/videos", params={"q": "plank"}).json()
    assert r["source"] == "link" and "plank" in r["search_url"]
    assert client.get("/api/videos", params={"q": "x"}).status_code == 422


def test_storage_survives_parallel_requests(tmp_path):
    """The page loads several videos at once; parallel reads used to crash SQLite."""
    from concurrent.futures import ThreadPoolExecutor

    from milo.storage import Storage

    store = Storage(tmp_path / "t.db")

    def work(i):
        q = f"exercise {i % 7}"
        store.save_video_cache(q, [{"id": str(i)}])
        assert store.get_video_cache(q) is not None
        store.get_conversation("coaching")
        return True

    with ThreadPoolExecutor(max_workers=12) as pool:
        assert all(pool.map(work, range(400)))
