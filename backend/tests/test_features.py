"""Nutrition, token usage, plan downloads and starting over."""
import io
import json

import pytest
from fastapi.testclient import TestClient

NUTRITION = {
    "summary": "Simple fuel for training days and rest days.",
    "principles": ["Eat a carb-rich meal 2-3 hours before runs", "Include protein at every meal"],
    "training_day": [{"meal": "Breakfast", "timing": "2-3 h before training", "suggestion": "Oats with banana"}],
    "rest_day": [{"meal": "Lunch", "timing": "", "suggestion": "Lentil salad"}],
    "hydration": "Drink regularly; a glass of water with every meal.",
    "notes": "Vegetarian. General guidance, not medical advice.",
}


@pytest.fixture
def client(fake_llm):
    import main
    from milo.coach import Coach
    from milo.storage import Storage

    main.store = Storage(":memory:")
    main.coach = Coach(main.store)
    return TestClient(main.app)


def events(response) -> list[dict]:
    return [json.loads(next(line[6:] for line in b.splitlines() if line.startswith("data: ")))
            for b in response.text.strip().split("\n\n")]


def say(client, text, mode="coaching", week=None):
    return events(client.post("/api/chat/stream", json={"mode": mode, "week": week, "text": text}))[-1]


def test_usage_is_reported_and_stored(client, fake_llm):
    fake_llm.replies = ["What's your goal?"]
    done = say(client, None)
    u = done["message"]["usage"]
    assert u["input_tokens"] > 0 and u["output_tokens"] > 0 and u["estimated"] is True
    assert client.get("/api/coaching").json()[-1]["usage"]["output_tokens"] == u["output_tokens"]


def test_cost_uses_model_prices():
    from milo.pricing import cost_usd

    assert cost_usd("github/gpt-4o-mini", 10_000, 10_000) == 0.0          # free tier
    assert cost_usd("openai/gpt-6-luna", 1_000_000, 1_000_000) == pytest.approx(0.60)


def test_nutrition_is_added_and_kept_on_updates(client, fake_llm, plan_json):
    fake_llm.replies = [f"Plan:\n```json\n{json.dumps(plan_json)}\n```\nWould you also like a nutrition plan?"]
    assert say(client, "make my plan")["plan_saved"]
    assert client.get("/api/plan").json()["plan"]["nutrition"] is None

    with_nutrition = {**plan_json, "nutrition": NUTRITION}
    fake_llm.replies = [f"Added.\n```json\n{json.dumps(with_nutrition)}\n```"]
    say(client, "Yes please")
    assert client.get("/api/plan").json()["plan"]["nutrition"]["training_day"][0]["meal"] == "Breakfast"

    # The model updates the training part and leaves out "nutrition": it is kept.
    fake_llm.replies = [f"Updated.\n```json\n{json.dumps(plan_json)}\n```"]
    say(client, "Make Tuesday 35 minutes")
    assert client.get("/api/plan").json()["plan"]["nutrition"]["summary"].startswith("Simple fuel")

    # An explicit null removes it.
    fake_llm.replies = [f"Removed.\n```json\n{json.dumps({**plan_json, 'nutrition': None})}\n```"]
    say(client, "Remove the nutrition plan")
    assert client.get("/api/plan").json()["plan"]["nutrition"] is None


def test_downloads(client, fake_llm, plan_json):
    assert client.get("/api/plan/export.pdf").status_code == 404
    fake_llm.replies = [f"```json\n{json.dumps({**plan_json, 'nutrition': NUTRITION})}\n```"]
    say(client, "plan please")

    pdf = client.get("/api/plan/export.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "attachment" in pdf.headers["content-disposition"]

    xlsx = client.get("/api/plan/export.xlsx")
    assert xlsx.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Plan", "Nutrition"]
    rows = list(wb["Plan"].iter_rows(values_only=True))
    header = rows.index(next(r for r in rows if r and r[0] == "Week"))
    assert len(rows) - header - 1 == 4 * 7                       # every day of 4 weeks
    assert rows[header][9] == "Done?"                             # tracking column
    assert client.get("/api/plan/export.docx").status_code == 422


def test_start_over_keeps_or_clears_plan(client, fake_llm, plan_json):
    fake_llm.replies = [f"```json\n{json.dumps(plan_json)}\n```"]
    say(client, "plan please")
    client.put("/api/profile", json={"name": "Alex", "age": 30})

    # Keep the plan: only the conversation is cleared.
    assert client.delete("/api/conversations/coaching").status_code == 200
    assert client.get("/api/coaching").json() == []
    assert client.get("/api/plan").json()["plan"] is not None

    # Clear the plan: plans and conversations go, profile stays.
    fake_llm.replies = ["Hi"]
    say(client, None)
    assert client.delete("/api/plan").status_code == 200
    assert client.get("/api/plan").json()["plan"] is None
    assert client.get("/api/coaching").json() == []
    assert client.get("/api/profile").json()["name"] == "Alex"
    assert client.delete("/api/conversations/checkin").status_code == 400


def test_compact_plan_is_accepted_and_sent_compact(client, fake_llm, plan_json):
    """Plans are written compactly (no rest days, no empty fields) so they generate faster."""
    from openpyxl import load_workbook

    from milo.coach import compact_plan
    from milo.schemas import TrainingPlan

    compact = json.loads(compact_plan(TrainingPlan.model_validate(plan_json)))
    assert all(s["type"] != "rest" for w in compact["weeks"] for s in w["sessions"])
    assert "distance_km\":null" not in compact_plan(TrainingPlan.model_validate(plan_json))
    assert len(compact_plan(TrainingPlan.model_validate(plan_json))) < len(json.dumps(plan_json, indent=2)) * 0.5

    # a compact plan from the model is a complete plan for the app
    fake_llm.replies = [f"Here it is.\n```json\n{json.dumps(compact)}\n```"]
    assert client.post("/api/coaching/message", json={"text": "plan"}).json()["plan_saved"]
    saved = client.get("/api/plan").json()["plan"]
    strength = next(x for x in saved["weeks"][0]["sessions"] if x["type"] == "strength")
    assert strength["distance_km"] is None and strength["target_intensity"]["kind"] in ("rpe", "none")   # defaults filled in
    ws = load_workbook(io.BytesIO(client.get("/api/plan/export.xlsx").content))["Plan"]
    assert sum(1 for r in ws.iter_rows(min_row=5) if r[1].value) == 7 * len(compact["weeks"])   # rest days back
    m = client.get("/api/weeks/1/metrics").json()["summary"]
    assert m["planned_sessions"] == len(compact["weeks"][0]["sessions"])

    # the model gets the compact plan as CURRENT_PLAN
    fake_llm.replies = ["Sure."]
    client.post("/api/coaching/message", json={"text": "anything"})
    system = fake_llm.calls[-1][0]["content"]
    assert '"type":"rest"' not in system and '"distance_km":null' not in system
