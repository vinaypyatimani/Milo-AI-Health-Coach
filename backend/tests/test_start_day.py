"""Starting a plan today (mid-week) or next Monday."""
import io
import json

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(fake_llm):
    import main
    from milo.coach import Coach
    from milo.storage import Storage

    main.store = Storage(":memory:")
    main.coach = Coach(main.store)
    return TestClient(main.app)


def save(client, fake_llm, plan, text="plan please"):
    fake_llm.replies = [f"Plan.\n```json\n{json.dumps(plan)}\n```"]
    assert client.post("/api/coaching/message", json={"text": text}).json()["plan_saved"]


def test_first_day_rules(plan_json):
    from milo.schemas import TrainingPlan

    p = TrainingPlan.model_validate({**plan_json, "first_day": "2026-10-15"})   # Thursday of week 1
    assert p.start_date.isoformat() == "2026-10-12" and p.begins.isoformat() == "2026-10-15"
    assert not p.is_active(1, "Tuesday") and p.is_active(1, "Thursday") and p.is_active(2, "Monday")
    assert TrainingPlan.model_validate({**plan_json, "first_day": "2026-10-12"}).first_day is None  # Monday start
    assert TrainingPlan.model_validate({**plan_json, "first_day": "2026-12-01"}).first_day is None  # not in week 1


def test_days_before_the_start_are_not_missed(client, fake_llm, plan_json):
    save(client, fake_llm, {**plan_json, "first_day": "2026-10-15"})
    # A run on Tuesday, before the plan started, must not count.
    client.post("/api/activities/manual", json={"date": "2026-10-13", "type": "run", "duration_min": 30})
    client.post("/api/activities/manual", json={"date": "2026-10-15", "type": "strength", "duration_min": 40})
    m = client.get("/api/weeks/1/metrics").json()
    assert [s["planned_day"] for s in m["sessions"]] == ["Thursday", "Saturday", "Sunday"]
    assert m["summary"]["missed"] == 2 and m["summary"]["extra_unplanned_activities"] == 0
    assert m["summary"]["actual_duration_min"] == 40
    week1 = client.get("/api/progress").json()[0]
    assert week1["planned_duration_min"] == 40 + 45 + 20


def test_update_keeps_first_day(client, fake_llm, plan_json):
    save(client, fake_llm, {**plan_json, "first_day": "2026-10-15"})
    changed = json.loads(json.dumps(plan_json))
    changed["weeks"][1]["focus"] = "Easier week"
    save(client, fake_llm, changed, "make week 2 easier")  # the model forgot first_day
    assert client.get("/api/plan").json()["plan"]["first_day"] == "2026-10-15"


def test_downloads_mark_days_before_start(client, fake_llm, plan_json):
    from openpyxl import load_workbook

    save(client, fake_llm, {**plan_json, "first_day": "2026-10-15"})
    ws = load_workbook(io.BytesIO(client.get("/api/plan/export.xlsx").content))["Plan"]
    sessions = [r[3].value for r in ws.iter_rows(min_row=5, max_row=11)]
    assert sessions[:3] == ["Before start"] * 3 and sessions[3] == "Full-body strength"
