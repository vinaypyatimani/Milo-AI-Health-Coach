"""Download the plan as a PDF (to read or print) or an Excel file (to track offline)."""
from __future__ import annotations

from io import BytesIO

from .schemas import DAYS, NutritionPlan, PlannedSession, TrainingPlan, UserProfile

ACCENT = "2F54EB"


def _target(s: PlannedSession) -> str:
    return s.target_intensity.value if s.target_intensity.kind != "none" else ""


def _exercises(s: PlannedSession) -> str:
    return "; ".join(f"{e.name} {e.detail}".strip() for e in s.exercises)


def _notes(s: PlannedSession) -> str:
    """Notes plus the exercise list, for the downloads."""
    ex = _exercises(s)
    return " — ".join(x for x in (s.notes, ex) if x)


def _sessions(plan: TrainingPlan):
    """Every day of every week, rest days included, in calendar order."""
    for w in plan.weeks:
        by_day = {s.day: s for s in w.sessions}
        for day in DAYS:
            s = by_day.get(day)
            yield w, day, plan.session_date(w.week_number, day), s


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def plan_to_xlsx(plan: TrainingPlan, profile: UserProfile | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=ACCENT)
    week_fill = PatternFill("solid", fgColor="EEF2FF")

    ws.append([plan.goal])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Starts {plan.begins.isoformat()}" + (f" · {profile.name}" if profile and profile.name else "")])
    ws.append([])
    headers = ["Week", "Day", "Date", "Session", "Type", "Planned (min)", "Distance (km)",
               "Target", "Notes", "Done?", "Actual (min)", "How it felt"]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical="center")

    for w, day, date, s in _sessions(plan):
        rest = s is None or s.type == "rest" or date < plan.begins
        ws.append([
            w.week_number, day, date,
            ("Before start" if date < plan.begins else "Rest") if rest else s.title, "" if rest else s.type,
            None if rest else s.duration_min, None if rest or s.distance_km is None else s.distance_km,
            "" if rest else _target(s), "" if rest else _notes(s),
            "" if rest else None, None, None,
        ])
        row = ws.max_row
        ws.cell(row, 3).number_format = "ddd d mmm"
        if w.week_number % 2 == 0:
            for c in range(1, 10):
                ws.cell(row, c).fill = week_fill
        if rest:
            for c in range(1, 10):
                ws.cell(row, c).font = Font(color="8A94A6")

    # "Done?" is a dropdown so the sheet works as an offline tracker.
    dv = DataValidation(type="list", formula1='"Yes,No,Moved"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J{header_row + 1}:J{ws.max_row}")

    widths = [7, 11, 13, 24, 10, 13, 13, 16, 60, 9, 12, 24]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:L{ws.max_row}"

    if plan.nutrition:
        _nutrition_sheet(wb.create_sheet("Nutrition"), plan.nutrition, head_font, head_fill)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nutrition_sheet(ws, n: NutritionPlan, head_font, head_fill) -> None:
    from openpyxl.styles import Alignment, Font

    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append(["Nutrition"])
    ws["A1"].font = Font(bold=True, size=14)
    if n.summary:
        ws.append([n.summary])
    ws.append([])
    if n.principles:
        ws.append(["Principles"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for p in n.principles:
            ws.append([f"• {p}"])
        ws.append([])
    for title, meals in (("Training day", n.training_day), ("Rest day", n.rest_day)):
        if not meals:
            continue
        ws.append([title, "When", "Suggestion"])
        for c in ws[ws.max_row]:
            c.font, c.fill = head_font, head_fill
        for m in meals:
            ws.append([m.meal, m.timing, m.suggestion])
            for c in ws[ws.max_row]:
                c.alignment = wrap
        ws.append([])
    if n.hydration:
        ws.append(["Hydration", "", n.hydration])
    if n.notes:
        ws.append(["Notes", "", n.notes])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 70


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def plan_to_pdf(plan: TrainingPlan, profile: UserProfile | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=plan.goal, author="Milo")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4,
                        textColor=colors.HexColor("#" + ACCENT))
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=8.5, leading=11)
    small = ParagraphStyle("s", parent=body, textColor=colors.HexColor("#6B7385"))

    def P(text: str, style=body) -> Paragraph:
        safe = (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    story = [P(plan.goal, h1),
             P(f"{len(plan.weeks)} weeks from {plan.begins.strftime('%d %b %Y')}"
               + (f" · {profile.name}" if profile and profile.name else "") + " · made with Milo", small)]

    for w in plan.weeks:
        story.append(P(f"Week {w.week_number}" + (f" — {w.focus}" if w.focus else ""), h2))
        rows = [[P(x, small) for x in ("Day", "Date", "Session", "Time", "Distance", "Target", "Notes", "Done")]]
        by_day = {s.day: s for s in w.sessions}
        for day in DAYS:
            s = by_day.get(day)
            d = plan.session_date(w.week_number, day).strftime("%d %b")
            if s is None or s.type == "rest" or plan.session_date(w.week_number, day) < plan.begins:
                label = "Before start" if plan.session_date(w.week_number, day) < plan.begins else "Rest"
                rows.append([P(day), P(d), P(label, small), "", "", "", "", ""])
            else:
                rows.append([P(day), P(d), P(s.title), P(f"{s.duration_min:g} min"),
                             P(f"{s.distance_km:g} km" if s.distance_km else ""), P(_target(s)), P(_notes(s)), "[  ]"])
        t = Table(rows, colWidths=[22 * mm, 16 * mm, 45 * mm, 18 * mm, 18 * mm, 28 * mm, 100 * mm, 12 * mm],
                  repeatRows=1)
        t.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#" + ACCENT)),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E3E7EF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ALIGN", (-1, 1), (-1, -1), "CENTER"),
        ]))
        story.append(t)

    n = plan.nutrition
    if n:
        story += [Spacer(1, 6 * mm), P("Nutrition", h1)]
        if n.summary:
            story.append(P(n.summary))
        for p in n.principles:
            story.append(P(f"• {p}"))
        for title, meals in (("Training day", n.training_day), ("Rest day", n.rest_day)):
            if meals:
                story.append(P(title, h2))
                t = Table([[P("Meal", small), P("When", small), P("Suggestion", small)]]
                          + [[P(m.meal), P(m.timing), P(m.suggestion)] for m in meals],
                          colWidths=[40 * mm, 45 * mm, 175 * mm])
                t.setStyle(TableStyle([
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#" + ACCENT)),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E3E7EF")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(t)
        if n.hydration:
            story += [P("Hydration", h2), P(n.hydration)]
        if n.notes:
            story += [Spacer(1, 3 * mm), P(n.notes, small)]

    story += [Spacer(1, 6 * mm),
              P("Milo is a coach, not a doctor. For pain, symptoms or medical conditions, see a professional.", small)]
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Analysis report (Analyze page)
# ---------------------------------------------------------------------------
def _inline_md(text: str) -> str:
    """Escape for reportlab and turn **bold** and *italic* into tags."""
    import re

    safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
    return re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", safe)


def report_to_pdf(stats: dict, report_md: str, profile: UserProfile | None = None) -> bytes:
    """Milo's written report plus the key numbers, as a printable PDF."""
    import re

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title="Training report", author="Milo")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12.5, spaceBefore=12, spaceAfter=4,
                        textColor=colors.HexColor("#" + ACCENT))
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=14)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#6B7385"))

    p = stats["period"]
    who = f" · {profile.name}" if profile and profile.name else ""
    story = [Paragraph("Training report", h1),
             Paragraph(f"{p['from']} to {p['to']} ({p['days']} days){who} · made with Milo", small)]

    # Milo's text: headings, bullets, numbered items, paragraphs
    items: list = []
    numbered = False

    def flush():
        nonlocal items
        if items:
            story.append(ListFlowable([ListItem(Paragraph(_inline_md(t), body), leftIndent=12) for t in items],
                                      bulletType="1" if numbered else "bullet", start="1" if numbered else None,
                                      leftIndent=14, bulletFontSize=9))
            items = []

    for line in report_md.splitlines():
        s = line.strip()
        bullet = re.match(r"^[-*•]\s+(.*)", s)
        num = re.match(r"^\d+[.)]\s+(.*)", s)
        if bullet or num:
            if items and numbered != bool(num):
                flush()
            numbered = bool(num)
            items.append((bullet or num).group(1))
            continue
        flush()
        if not s:
            continue
        if s.startswith("#"):
            story.append(Paragraph(_inline_md(s.lstrip("#").strip()), h2))
        else:
            story.append(Paragraph(_inline_md(s), body))
    flush()

    # Key numbers, straight from the analysis (not from the model's text)
    t = stats.get("totals") or {}
    rows = [["Sessions", str(t.get("sessions", "–"))], ["Hours", str(t.get("hours", "–"))],
            ["Sessions per week", str(t.get("sessions_per_week", "–"))],
            ["Distance", f"{t.get('distance_km', 0)} km"]]
    if stats.get("intensity"):
        i = stats["intensity"]
        rows.append(["Easy / moderate / hard", f"{i['easy_pct']}% / {i['moderate_pct']}% / {i['hard_pct']}%"])
    if stats.get("recovery") and stats["recovery"].get("avg_sleep_hours"):
        r = stats["recovery"]
        rows.append(["Average sleep", f"{r['avg_sleep_hours']} h"])
        if r.get("resting_hr_avg"):
            rows.append(["Resting heart rate", f"{r['resting_hr_avg']} bpm"])
    story += [Paragraph("Key numbers", h2), _kv_table(rows, Table, TableStyle, colors, mm)]

    weekly = stats.get("weekly") or []
    if weekly:
        wrows = [["Week of", "Sessions", "Minutes", "Distance (km)"]] + [
            [w["week_start"], str(w["sessions"]), str(w["minutes"]), str(w["distance_km"] or "")] for w in weekly]
        wt = Table(wrows, colWidths=[40 * mm, 30 * mm, 30 * mm, 35 * mm], repeatRows=1)
        wt.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 9), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6B7385")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#" + ACCENT)),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E3E7EF")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ]))
        story += [Paragraph("Week by week", h2), wt]

    story += [Spacer(1, 6 * mm),
              Paragraph("Milo is a coach, not a doctor. For pain, symptoms or medical conditions, see a professional.",
                        small)]
    doc.build(story)
    return buf.getvalue()


def _kv_table(rows, Table, TableStyle, colors, mm):
    t = Table(rows, colWidths=[60 * mm, 60 * mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#6B7385")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E3E7EF")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t
